from concurrent.futures import ThreadPoolExecutor
import io
import pytest
from docx import Document
from fastapi.testclient import TestClient
from sqlalchemy import select
from testpilot.api import app
from testpilot.config import settings
from testpilot.documents import MAX_BYTES
from testpilot.storage import Record, session, read_record
from testpilot.workspace import project_scope
from testpilot import sources
from test_chat_documents import pdf_bytes


@pytest.fixture
def library(tmp_path, monkeypatch):
    monkeypatch.setattr(settings, "data", tmp_path)
    client = TestClient(app, base_url="https://testserver", headers={"Authorization": "Bearer " + settings.access_token})
    projects = [client.post("/api/projects", json={"name": "项目资料测试 " + name}).json()["id"] for name in ["A", "B"]]
    yield client, projects
    with session() as db:
        for row in db.scalars(select(Record).where(Record.project_id.in_(projects) | Record.id.in_(projects))):
            db.delete(row)


def upload(client, project, name="需求.md", content=b"Requirement text"):
    return client.post("/api/sources/upload", params={"filename": name}, headers={"X-TestPilot-Project": project}, content=content)


def tool(client, project, name, arguments):
    return client.post("/api/tools/" + name, headers={"X-TestPilot-Project": project}, json={"arguments": arguments})


def test_file_persistence_lazy_reads_search_and_isolation(library):
    client, (a, b) = library
    text = "需求开始\n" + "status=OK\n" * 1000 + "UNIQUE-END-MARKER"
    response = upload(client, a, content=text.encode())
    assert response.status_code == 201, response.text
    source = response.json()
    assert source["storage"] == "file" and source["readable"]
    assert source["characters"] == len(text) and "content" not in source
    headers = {"X-TestPilot-Project": a}
    state = client.get("/api/state", headers=headers).json()
    assert state["sources"] == [source]
    assert "UNIQUE-END-MARKER" not in str(state["sources"])
    found = tool(client, a, "sources_search", {"query": "unique-end-marker"}).json()
    assert [s["id"] for s in found] == [source["id"]]
    page = tool(client, a, "sources_read", {"id": source["id"], "limit": 80}).json()
    assert page["content"] == text[:80] and page["nextOffset"] == 80
    assert client.get(f"/api/sources/{source['id']}/text?offset=80&limit=80", headers=headers).json()["content"] == text[80:160]
    assert client.get(source["downloadUrl"]).content == text.encode()
    assert TestClient(app, base_url="https://testserver").get(source["downloadUrl"]).status_code == 401
    assert client.get(f"/api/sources/{source['id']}/text", headers={"X-TestPilot-Project": b}).status_code == 404
    assert client.get(source["downloadUrl"], headers={"X-TestPilot-Project": b}).status_code == 404
    assert tool(client, b, "sources_search", {"query": ""}).json() == []
    assert client.get(f"/api/sources/{source['id']}/text?limit=16001", headers=headers).status_code == 422
    context = tool(client, a, "project_context", {}).json()
    assert context["sourceCount"] == 1 and context["sources"][0]["id"] == source["id"]
    assert "UNIQUE-END-MARKER" not in str(context)


def test_archive_restore_preserves_original_but_excludes_agent_reads(library):
    client, (a, b) = library
    source = upload(client, a).json(); headers = {"X-TestPilot-Project": a}
    url = f"/api/sources/{source['id']}/archive"
    assert client.post(url, headers={"X-TestPilot-Project": b}, json={"archived": True}).status_code == 404
    assert client.post(url, headers=headers, json={"archived": True}).json()["archived"]
    assert tool(client, a, "sources_search", {"query": ""}).json() == []
    assert tool(client, a, "project_context", {}).json()["sourceCount"] == 0
    assert tool(client, a, "sources_read", {"id": source["id"]}).status_code == 400
    assert client.get(source["downloadUrl"]).content == b"Requirement text"
    assert not client.post(url, headers=headers, json={"archived": False}).json()["archived"]
    assert tool(client, a, "sources_read", {"id": source["id"]}).json()["content"] == "Requirement text"


@pytest.mark.parametrize("name,content,reason", [("scan.pdf", pdf_bytes(), "OCR"), ("broken.docx", b"not a zip", "损坏"), ("assets.zip", b"opaque\x00bytes", "只保存原文件")])
def test_unreadable_files_still_download_with_explicit_reason(library, name, content, reason):
    client, (a, b) = library
    response = upload(client, a, name, content)
    assert response.status_code == 201, response.text
    source = response.json()
    assert not source["readable"] and reason in str(source["warnings"])
    assert client.get(source["downloadUrl"]).content == content
    result = tool(client, a, "sources_read", {"id": source["id"]}).json()
    assert not result["readable"] and result["content"] == "" and result["nextOffset"] is None
    assert tool(client, a, "sources_search", {"query": name}).json()[0]["id"] == source["id"]


def test_docx_pdf_text_and_legacy_materials_share_tools(library):
    client, (a, b) = library
    doc = Document(); doc.add_paragraph("需求说明")
    table = doc.add_table(rows=1, cols=2); table.cell(0, 0).text = "状态"; table.cell(0, 1).text = "OK"
    buffer = io.BytesIO(); doc.save(buffer)
    uploaded = upload(client, a, "需求.docx", buffer.getvalue()).json()
    assert "状态 | OK" in tool(client, a, "sources_read", {"id": uploaded["id"]}).json()["content"]
    pdf = upload(client, a, "requirements.pdf", pdf_bytes("HTTP 200 PDF-PROJECT-OK")).json()
    assert pdf["readable"] and pdf["pages"] == 1
    legacy = client.post("/api/sources", headers={"X-TestPilot-Project": a}, json={"name": "接口说明", "content": "必须拒绝不存在的账号"}).json()
    assert legacy["storage"] == "text" and "content" not in legacy
    assert client.get(legacy["downloadUrl"]).content.decode() == "必须拒绝不存在的账号"
    assert len(tool(client, a, "sources_search", {"query": "", "limit": 1, "offset": 1}).json()) == 1
    # 检查聊天附件仍按所属 Task 隔离，不因项目相同就自动加入共享资料库；只有明确上传为项目资料的文件才可供其他对话读取。
    assert client.post("/api/attachments", params={"filename": "附件.txt"}, headers={"X-TestPilot-Project": a}, content=b"chat only").status_code == 201
    assert len(client.get("/api/sources", headers={"X-TestPilot-Project": a}).json()) == 3


def test_download_always_attachment_and_symlinks_cannot_cross_file_boundaries(library, tmp_path):
    client, (a, b) = library
    source = upload(client, a, "page.html", b"<script>alert(1)</script>").json()
    response = client.get(source["downloadUrl"])
    assert response.headers["content-type"] == "application/octet-stream"
    assert response.headers["x-content-type-options"] == "nosniff"
    assert response.headers["content-disposition"].startswith("attachment;")
    with project_scope(a): record = read_record(source["id"], "source")
    file = sources.file_path(record)
    other = tmp_path / "private.txt"; other.write_text("OTHER FILE")
    file.unlink(); file.symlink_to(other)
    assert client.get(source["downloadUrl"]).status_code == 404


def test_upload_limits_cleanup_and_archived_storage_quota(library, monkeypatch):
    client, (a, b) = library
    assert upload(client, a, "../bad.txt").status_code == 400
    assert upload(client, a, content=b"").status_code == 400
    assert upload(client, a, content=b"x" * (MAX_BYTES + 1)).status_code == 413
    monkeypatch.setattr(sources, "MAX_PROJECT_BYTES", 20)
    source = upload(client, a, "first.zip", b"1" * 15).json()
    client.post(f"/api/sources/{source['id']}/archive", headers={"X-TestPilot-Project": a}, json={"archived": True})
    assert upload(client, a, "second.zip", b"2" * 6).status_code == 400
    assert len(list((settings.data / "sources").iterdir())) == 1


def test_concurrent_upload_quota_is_atomic(library, monkeypatch):
    client, (a, b) = library
    monkeypatch.setattr(sources, "MAX_FILES", 1)
    def attempt(number):
        with project_scope(a):
            try: return sources.upload_source(f"file{number}.zip", b"data")["id"]
            except ValueError as error: return str(error)
    with ThreadPoolExecutor(max_workers=2) as pool:
        results = list(pool.map(attempt, [1, 2]))
    assert sum(value.startswith("SRC-") for value in results) == 1
    assert len(client.get("/api/sources", headers={"X-TestPilot-Project": a}).json()) == 1
    assert len(list((settings.data / "sources").iterdir())) == 1
