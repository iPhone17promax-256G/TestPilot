import io
import json
from datetime import timedelta
from pathlib import Path
import pytest
from docx import Document
from pypdf import PdfWriter
from pypdf.generic import NameObject, DictionaryObject, DecodedStreamObject
from fastapi.testclient import TestClient
from sqlalchemy import select
from testpilot.api import app
from testpilot.config import settings
from testpilot.documents import extract_document, validate_name, MAX_BYTES
from testpilot.attachments import message_content
from testpilot.job_queue import enqueue
from queue_helpers import mark_running
from testpilot.storage import Record, session, uid, now, read_record
from testpilot.workspace import project_scope
from testpilot import testing as t

def pdf_bytes(text=None, encrypted=False):
    writer = PdfWriter()
    page = writer.add_blank_page(width=300, height=300)
    if text:
        font=DictionaryObject({NameObject("/Type"):NameObject("/Font"),NameObject("/Subtype"):NameObject("/Type1"),NameObject("/BaseFont"):NameObject("/Helvetica")})
        page[NameObject("/Resources")]=DictionaryObject({NameObject("/Font"):DictionaryObject({NameObject("/F1"):font})})
        stream=DecodedStreamObject()
        stream.set_data(f"BT /F1 12 Tf 20 200 Td ({text}) Tj ET".encode())
        page[NameObject("/Contents")]=writer._add_object(stream)
    if encrypted:writer.encrypt("test-password")
    buffer=io.BytesIO();writer.write(buffer)
    return buffer.getvalue()

@pytest.mark.parametrize("encoding",["utf-8-sig","utf-16","gb18030"])
def test_text_encodings(tmp_path,encoding):
    path=tmp_path/"需求.txt";path.write_bytes("接口需求：状态为 OK".encode(encoding))
    assert extract_document(path)["text"]=="接口需求：状态为 OK"

def test_docx_preserves_table_position(tmp_path):
    doc=Document();doc.add_paragraph("需求开始")
    table=doc.add_table(rows=1,cols=2);table.cell(0,0).text="接口路径";table.cell(0,1).text="/api/v4/system/ping"
    doc.add_paragraph("状态必须为 OK")
    path=tmp_path/"需求.docx";doc.save(path)
    text=extract_document(path)["text"]
    assert text.index("需求开始")<text.index("接口路径 | /api/v4/system/ping")<text.index("状态必须为 OK")

def test_pdf_text_and_scan_error(tmp_path):
    path=tmp_path/"requirements.pdf";path.write_bytes(pdf_bytes("HTTP 200 and JSON status OK"))
    assert "HTTP 200 and JSON status OK" in extract_document(path)["text"]
    path.write_bytes(pdf_bytes())
    with pytest.raises(ValueError,match="OCR"):extract_document(path)

def test_encrypted_pdf_and_broken_docx(tmp_path):
    path=tmp_path/"locked.pdf";path.write_bytes(pdf_bytes("test",encrypted=True))
    with pytest.raises(ValueError,match="加密"):extract_document(path)
    with pytest.raises(ValueError):validate_name("../secret.txt")
    with pytest.raises(ValueError):validate_name("legacy.doc")

@pytest.fixture
def workspace(tmp_path,monkeypatch):
    monkeypatch.setattr(settings,"data",tmp_path)
    client=TestClient(app, base_url="https://testserver",headers={"Authorization":"Bearer "+settings.access_token})
    projects=[client.post("/api/projects",json={"name":"附件测试 "+x}).json()["id"] for x in ["A","B"]]
    yield client,projects
    with session() as db:
        for row in db.scalars(select(Record).where((Record.project_id.in_(projects)) | (Record.id.in_(projects)))):
            db.delete(row)

def upload(client,project,name="需求.md",content="# 需求\nHTTP 200；status 为 OK".encode()):
    return client.post("/api/attachments",params={"filename":name},headers={"X-TestPilot-Project":project},content=content)

def test_upload_preview_download_scope_and_paging(workspace):
    client,(a,b)=workspace
    text="业务要求\n"+"status=OK\n"*1000
    response=upload(client,a,content=text.encode())
    assert response.status_code==201,response.text
    attachment=response.json();identity=attachment["id"]
    assert attachment["characters"]==len(text.strip()) and "textPath" not in attachment
    first=client.get(f"/api/attachments/{identity}/text?limit=80",headers={"X-TestPilot-Project":a}).json()
    assert first["content"]==text[:80] and first["nextOffset"]==80
    second=client.get(f"/api/attachments/{identity}/text?offset=80&limit=80",headers={"X-TestPilot-Project":a}).json()
    assert second["content"]==text[80:160]
    assert client.get(attachment["downloadUrl"]).content==text.encode()
    assert client.get(f"/api/attachments/{identity}/text",headers={"X-TestPilot-Project":b}).status_code==404
    assert TestClient(app, base_url="https://testserver").get(attachment["downloadUrl"]).status_code==401
    assert client.get(f"/api/attachments/{identity}/text?limit=16001",headers={"X-TestPilot-Project":a}).status_code==422

def test_followup_attachment_persists_and_tool_is_conversation_bound(workspace):
    client,(a,b)=workspace
    file=upload(client,a).json();other=upload(client,a,name="另一个.txt",content=b"unrelated").json()
    identity=uid("TASK")
    with project_scope(a),session() as db:
        db.add(Record(id=identity,kind="task",status="running",payload=dict(id=identity,status="running",messages=[],events=[],runIds=[],caseIds=[],goal="原任务",createdAt=t.stamp())))
        enqueue(db,db.get(Record,identity),"agent")
    mark_running(identity)
    headers={"X-TestPilot-Project":a}
    for _ in range(2):
        response=client.post(f"/api/tasks/{identity}/messages",headers=headers,json={"content":"按附件设计，别执行","attachment_ids":[file["id"]],"key":"same-document"})
        assert response.status_code==202,response.text
    with project_scope(a):
        message=read_record(identity)["messages"][0]
        assert len(read_record(identity)["messages"])==1
        assert message["attachments"][0]["id"]==file["id"]
        assert file["id"] in message_content(message)
        assert "HTTP 200" not in message_content(message) # 检查初始 Agent 上下文只包含附件索引而不含整份正文；正文应由 attachments_read 分页读取，避免需求文件自动占满每轮提示词。
        assert t.take_user_messages(identity)[0]["attachments"][0]["id"]==file["id"]
    result=client.post("/api/tools/attachments_read",headers=headers,json={"task_id":identity,"arguments":{"id":file["id"]}})
    assert result.status_code==200 and "status 为 OK" in result.json()["content"]
    denied=client.post("/api/tools/attachments_read",headers=headers,json={"task_id":identity,"arguments":{"id":other["id"]}})
    assert denied.status_code==400
    foreign=upload(client,b).json()
    assert client.post(f"/api/tasks/{identity}/messages",headers=headers,json={"content":"需求","attachment_ids":[foreign["id"]]}).status_code==404

def test_attachment_only_message_and_rejections(workspace):
    client,(a,b)=workspace
    assert upload(client,a,name="bad.exe").status_code==400
    assert upload(client,a,content=b"").status_code==400
    assert upload(client,a,content=b"x"*(MAX_BYTES+1)).status_code==413
    assert upload(client,a,name="scan.pdf",content=pdf_bytes()).status_code==400
    assert upload(client,a,name="broken.docx",content=b"not a zip").status_code==400
    file=upload(client,a).json()
    from testpilot.attachments import attachment_refs
    with project_scope(a):
        with pytest.raises(ValueError):attachment_refs([file["id"]]*2)
        with pytest.raises(ValueError):t.task_message("missing","")
        # 这里仅验证首条消息和附件引用的 Task 构造结果；用替身避免提交真实 Agent 作业，不调用模型供应商、不产生模型费用。
        env=t.environment_save("测试环境","http://target.test")
        collected=[]
        class FakeDB:
            info={}
            def add(self,row):collected.append(row)
        original=t.mutate
        t.mutate=lambda scope,key,args,action:action(FakeDB())
        try:
            task=t.task_start("",env["id"],attachment_ids=[file["id"]])
            assert "请阅读附件" in task["goal"]
            assert task["messages"][0]["attachments"][0]["id"]==file["id"]
        finally:t.mutate=original


def test_pdf_upload_through_bounded_parser(workspace):
    client,(a,b)=workspace
    response=upload(client,a,name="requirements.pdf",content=pdf_bytes("Requirement marker PDF-OK-20261006"))
    assert response.status_code==201,response.text
    attachment=response.json()
    assert attachment["pages"]==1
    text=client.get(f"/api/attachments/{attachment['id']}/text",headers={"X-TestPilot-Project":a}).json()
    assert "PDF-OK-20261006" in text["content"]
